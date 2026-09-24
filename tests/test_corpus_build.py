"""The clean baseline must genuinely be clean, and reproducible."""

from docx import Document
from docx.oxml.ns import qn

from lfe.corpus import build, make_contract
from lfe.corpus.model import BlockKind
from lfe.text import stored_paragraphs


def test_contract_is_deterministic() -> None:
    assert make_contract(7) == make_contract(7)
    assert make_contract(7) != make_contract(8)


def test_build_is_byte_identical_across_runs(tmp_path) -> None:
    contract = make_contract(3)
    first = build(contract, tmp_path / "a.docx").read_bytes()
    second = build(contract, tmp_path / "b.docx").read_bytes()
    assert first == second


def test_every_block_reaches_the_document(tmp_path) -> None:
    contract = make_contract(1)
    path = build(contract, tmp_path / "c.docx")

    stored = stored_paragraphs(path)
    assert stored == [block.text for block in contract.blocks]


def test_clause_numbers_are_real_numbering_not_typed_text(tmp_path) -> None:
    """The whole point of the clean baseline: numbers come from numbering.xml."""
    contract = make_contract(1)
    path = build(contract, tmp_path / "d.docx")

    numbered = [b for b in contract.blocks if b.number is not None]
    assert numbered, "fixture should contain numbered clauses"

    for line in stored_paragraphs(path):
        for block in numbered:
            assert not line.startswith(f"{block.number} ")


def test_numbered_paragraphs_carry_numPr_at_the_right_level(tmp_path) -> None:
    contract = make_contract(1)
    path = build(contract, tmp_path / "e.docx")

    document = Document(str(path))
    paragraphs = document.element.body.findall(qn("w:p"))
    assert len(paragraphs) == len(contract.blocks)

    for element, block in zip(paragraphs, contract.blocks, strict=True):
        num_pr = element.find(qn("w:pPr") + "/" + qn("w:numPr"))
        if block.number is None:
            assert num_pr is None
        else:
            assert num_pr is not None
            ilvl = num_pr.find(qn("w:ilvl")).get(qn("w:val"))
            assert int(ilvl) == block.level - 1


def test_headings_use_template_styles(tmp_path) -> None:
    contract = make_contract(2)
    path = build(contract, tmp_path / "f.docx")

    document = Document(str(path))
    heading_levels = [b.level for b in contract.blocks if b.kind is BlockKind.HEADING]
    used = [
        p.style.name
        for p, b in zip(document.paragraphs, contract.blocks, strict=True)
        if b.kind is BlockKind.HEADING
    ]
    assert used == [f"Heading {level}" for level in heading_levels]

"""Render a `Contract` to a clean, well-formed .docx.

"Clean" means what the normaliser is aiming at: every clause carries a real
style and real list numbering, and nothing is achieved by direct formatting.
The defect injectors take these files and make them ugly in specific,
labelled ways.

Output is deterministic — the same contract always produces the same bytes —
so golden files can be committed and diffed.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentT

from . import numbering
from .model import Block, BlockKind, Contract

# Pinned so repeated runs are byte-identical.
EPOCH = dt.datetime(2020, 1, 1, 0, 0, 0, tzinfo=dt.UTC)
ZIP_DATE_TIME = (2020, 1, 1, 0, 0, 0)

HEADING_STYLES = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}


def _add_block(document: DocumentT, block: Block, num_id: int) -> None:
    if block.kind is BlockKind.TITLE:
        document.add_paragraph(block.text, style="Title")
        return

    if block.kind is BlockKind.SCHEDULE_TITLE:
        document.add_paragraph(block.text, style="Heading 1")
        return

    if block.kind is BlockKind.HEADING:
        style = HEADING_STYLES[min(block.level, 4)]
        paragraph = document.add_paragraph(block.text, style=style)
        if block.number is not None:
            numbering.apply(paragraph, num_id, block.level)
        return

    if block.kind is BlockKind.DEFINITION:
        document.add_paragraph(block.text, style="Normal")
        return

    paragraph = document.add_paragraph(block.text, style="Normal")
    if block.number is not None:
        numbering.apply(paragraph, num_id, block.level)


def build(contract: Contract, path: str | Path) -> Path:
    """Write `contract` to `path` as a clean .docx and return the path."""
    document = Document()
    num_id = numbering.install(document)

    for block in contract.blocks:
        _add_block(document, block, num_id)

    core = document.core_properties
    core.title = contract.name
    core.author = "Legal Format Engine corpus generator"
    core.created = EPOCH
    core.modified = EPOCH
    core.last_modified_by = "corpus"
    core.revision = 1

    buffer = io.BytesIO()
    document.save(buffer)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_repack_deterministically(buffer.getvalue()))
    return path


def _repack_deterministically(blob: bytes) -> bytes:
    """Rewrite a .docx zip with pinned entry timestamps.

    python-docx saves through `ZipFile.writestr`, which stamps every member with
    the current clock at 2-second resolution. That makes otherwise identical
    builds differ in their bytes depending on when they ran, which would make
    golden files in the corpus churn for no reason. Member order and contents
    are preserved exactly; only the timestamps are replaced.
    """
    source = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            pinned = zipfile.ZipInfo(item.filename, date_time=ZIP_DATE_TIME)
            pinned.compress_type = item.compress_type
            pinned.external_attr = item.external_attr
            target.writestr(pinned, source.read(item.filename))
    return out.getvalue()
